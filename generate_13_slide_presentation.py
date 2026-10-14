"""
generate_13_slide_presentation.py
Official 13-slide M.Tech Research Seminar Presentation for MSBA-Net.
Strictly follows academic presentation flow: Bullets for explanatory slides, native charts/tables for results,
and only Slide 5 contains a clean, editable architecture diagram.
Output: MSBA-Net_Amen_Parmar_Presentation.pptx
"""

import os, sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_TICK_MARK, XL_LABEL_POSITION

# -------------------------------------------------------------
# CONSTANTS & PALETTE
# -------------------------------------------------------------
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

FONT_PRIMARY = "Segoe UI"
FONT_HEADING = "Segoe UI"

COLOR_BG = RGBColor(248, 247, 243)         # #F8F7F3 Off-white
COLOR_NAVY = RGBColor(11, 25, 44)          # #0B192C Deep Navy
COLOR_NAVY_LIGHT = RGBColor(30, 41, 59)    # #1E293B Dark Slate
COLOR_SLATE = RGBColor(71, 85, 105)        # #475569 Secondary Slate
COLOR_MUTED = RGBColor(148, 163, 184)      # #94A3B8 Muted Line
COLOR_BORDER = RGBColor(15, 23, 42)        # #0F172A Crisp Card Border
COLOR_CARD_WHITE = RGBColor(255, 255, 255) # Pure White
COLOR_CARD_ALT = RGBColor(246, 248, 250)   # Clean Ice Gray

# Accents
COLOR_AMBER = RGBColor(217, 119, 6)        # #D97706 Muted Amber / Highlight
COLOR_AMBER_BG = RGBColor(254, 243, 199)   # #FEF3C7 Soft Amber
COLOR_GREEN = RGBColor(5, 150, 105)        # #059669 Emerald / Results
COLOR_GREEN_BG = RGBColor(209, 250, 229)   # #D1FAE5 Soft Green
COLOR_BLUE = RGBColor(37, 99, 235)         # #2563EB Blue
COLOR_BLUE_BG = RGBColor(219, 234, 254)    # #DBEAFE Soft Blue
COLOR_PURPLE = RGBColor(124, 58, 237)      # #7C3AED BGFF Purple
COLOR_PURPLE_BG = RGBColor(243, 232, 255)  # #F3E8FF Soft Purple
COLOR_RED = RGBColor(220, 38, 38)          # #DC2626 Failure/Degradation
COLOR_RED_BG = RGBColor(254, 226, 226)     # #FEE2E2 Soft Red

GRID_IMAGE_PATH = os.path.join(os.path.dirname(__file__), "scratch", "notebooklm_bg.png")
SLIDE3_DEFECT_PATH = os.path.join(os.path.dirname(__file__), "scratch", "slide3_defect_sample.png")
SLIDE9_QUAL_PATH = os.path.join(os.path.dirname(__file__), "scratch", "slide9_qualitative_crop.png")

# -------------------------------------------------------------
# HELPER FUNCTIONS
# -------------------------------------------------------------
def set_slide_background(slide):
    """Sets subtle technical grid background."""
    if os.path.exists(GRID_IMAGE_PATH):
        slide.shapes.add_picture(GRID_IMAGE_PATH, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT)
    else:
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

def add_clean_paragraph(text_frame, text, font_name=FONT_PRIMARY, font_size=Pt(9.0), bold=False, color=COLOR_SLATE, alignment=PP_ALIGN.LEFT, space_after=Pt(3)):
    """Safely adds paragraph(s) to a text frame with explicit font properties and ensures word wrap."""
    text_frame.word_wrap = True
    lines = str(text).split("\n")
    p = None
    for idx, line in enumerate(lines):
        if len(text_frame.paragraphs) == 1 and text_frame.paragraphs[0].text == "":
            p = text_frame.paragraphs[0]
        else:
            p = text_frame.add_paragraph()
        p.text = line
        p.font.name = font_name
        p.font.size = font_size
        p.font.bold = bold
        p.font.color.rgb = color
        p.alignment = alignment
        p.space_after = space_after if idx == len(lines) - 1 else Pt(1)
    return p

def add_header(slide, category_tag, title_text, subtitle_text):
    """Adds a standard NotebookLM-styled academic header."""
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.73), Inches(1.10))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    add_clean_paragraph(tf, category_tag.upper(), font_size=Pt(9.0), bold=True, color=COLOR_AMBER, space_after=Pt(2))
    add_clean_paragraph(tf, title_text, font_name=FONT_HEADING, font_size=Pt(22), bold=True, color=COLOR_NAVY, space_after=Pt(2))
    add_clean_paragraph(tf, subtitle_text, font_size=Pt(11.0), bold=False, color=COLOR_SLATE, space_after=Pt(0))

def add_card(slide, x, y, w, h, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.2), shape_type=MSO_SHAPE.ROUNDED_RECTANGLE):
    """Draws a clean, outlined card shape (editable in PPTX)."""
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width
    else:
        shape.line.fill.background()
    return shape

def add_pill_badge(slide, x, y, w, h, text, text_color=COLOR_NAVY, bg_color=COLOR_CARD_ALT, border_color=COLOR_BORDER, font_size=Pt(8.5)):
    """Draws a clean pill-shaped label badge."""
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    pill.fill.solid()
    pill.fill.fore_color.rgb = bg_color
    if border_color:
        pill.line.color.rgb = border_color
        pill.line.width = Pt(1.0)
    else:
        pill.line.fill.background()
    tf = pill.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_PRIMARY
    p.font.size = font_size
    p.font.bold = True
    p.font.color.rgb = text_color
    p.alignment = PP_ALIGN.CENTER
    return pill

def set_speaker_note(slide, note_text):
    """Adds natural spoken script to the slide speaker notes."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = note_text.strip()


# =============================================================
# SLIDE 1: TITLE SLIDE
# =============================================================
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    start_x = Inches(0.8)

    # Left Container: Title and Author Metadata
    w_left = Inches(7.6)
    y_card = Inches(1.2)
    h_card = Inches(5.3)
    add_card(slide, start_x, y_card, w_left, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))

    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.4), Inches(3.2), Inches(0.32), "M.TECH (AIML) RESEARCH SEMINAR", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(8.5))

    tb_title = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.85), w_left - Inches(0.8), Inches(2.2))
    tf_t = tb_title.text_frame
    add_clean_paragraph(tf_t, "MSBA-Net: Multi-Scale Feature Enhancement and Boundary-Aware Learning for Small and Low-Contrast Steel Surface Defect Segmentation", font_name=FONT_HEADING, font_size=Pt(21), bold=True, color=COLOR_NAVY, space_after=Pt(8))
    add_clean_paragraph(tf_t, "Deep Learning-Based Surface Defect Segmentation in Metallurgical Manufacturing", font_size=Pt(12), bold=False, color=COLOR_SLATE, space_after=Pt(0))

    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, start_x + Inches(0.4), y_card + Inches(3.3), w_left - Inches(0.8), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_MUTED
    line.line.fill.background()

    # Author & Presenter Details
    tb_meta = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(3.45), w_left - Inches(0.8), Inches(1.5))
    tf_m = tb_meta.text_frame
    add_clean_paragraph(tf_m, "Presenter:", font_size=Pt(9.0), bold=True, color=COLOR_AMBER, space_after=Pt(1))
    add_clean_paragraph(tf_m, "Amen Parmar  •  M.Tech (AIML)  •  Roll No: 2502010010", font_size=Pt(11.0), bold=True, color=COLOR_NAVY, space_after=Pt(6))
    add_clean_paragraph(tf_m, "Co-authors: Amit Kumar Bairwa, Varun Kumar", font_size=Pt(9.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(3))
    add_clean_paragraph(tf_m, "Department of Computer Science and Engineering  |  Manipal University Jaipur, Jaipur, Rajasthan, India", font_size=Pt(9.0), bold=False, color=COLOR_SLATE, space_after=Pt(0))

    # Right Container: Minimal Academic Visual Card
    x_right = start_x + w_left + Inches(0.35)
    w_right = Inches(3.78)
    add_card(slide, x_right, y_card, w_right, h_card, bg_color=COLOR_CARD_ALT, border_color=COLOR_BORDER, border_width=Pt(1.5))

    add_pill_badge(slide, x_right + Inches(0.3), y_card + Inches(0.4), Inches(2.6), Inches(0.30), "CONCEPTUAL WORKFLOW", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.0))

    tb_rv = slide.shapes.add_textbox(x_right + Inches(0.3), y_card + Inches(0.9), w_right - Inches(0.6), Inches(4.0))
    tf_rv = tb_rv.text_frame
    add_clean_paragraph(tf_rv, "Industrial Objective", font_size=Pt(12.0), bold=True, color=COLOR_NAVY, space_after=Pt(6))
    add_clean_paragraph(tf_rv, "Automated visual inspection on high-speed steel rolling lines demands pixel-precise segmentation of micro-scale and low-contrast surface flaws.", font_size=Pt(9.5), bold=False, color=COLOR_SLATE, space_after=Pt(14))

    add_clean_paragraph(tf_rv, "Key Innovations in MSBA-Net:", font_size=Pt(10.5), bold=True, color=COLOR_NAVY, space_after=Pt(4))
    add_clean_paragraph(tf_rv, "1. MSFP: Preserves micro-scale features across scales without pooling loss.", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(6))
    add_clean_paragraph(tf_rv, "2. BAS: Auxiliary stream capturing 1-pixel morphological defect boundaries.", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(6))
    add_clean_paragraph(tf_rv, "3. BGFF: Residual spatial attention gating guiding decoder feature fusion.", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(14))

    add_clean_paragraph(tf_rv, "Validated on NEU Surface Defect Database", font_size=Pt(8.5), bold=True, color=COLOR_GREEN, alignment=PP_ALIGN.CENTER, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 1 — TITLE (45–55 seconds)
Speaker Script:
"Respected committee members, examiners, and faculty: Good morning.

My name is Amen Parmar, M.Tech in Artificial Intelligence and Machine Learning at Manipal University Jaipur, Roll Number 2502010010. Today, I am presenting our research paper entitled: 'MSBA-Net: Multi-Scale Feature Enhancement and Boundary-Aware Learning for Small and Low-Contrast Steel Surface Defect Segmentation', co-authored with Dr. Amit Kumar Bairwa and Dr. Varun Kumar.

In metallurgical manufacturing, automated surface defect inspection of flat-rolled steel is an indispensable foundation of quality assurance. However, traditional computer vision and standard deep learning models face severe failure modes when confronted with microscopic flaws and weak, low-contrast boundaries. In this presentation, I will explain the objectives, methodology, benchmark performance, and operational boundaries of MSBA-Net."
""")


# =============================================================
# SLIDE 2: OBJECTIVE
# =============================================================
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Objective", "Core research goals and technical aims of the proposed framework.")

    start_x = Inches(0.8)
    w_main = Inches(11.73)
    y_card = Inches(1.65)
    h_card = Inches(5.25)
    add_card(slide, start_x, y_card, w_main, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))

    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.2), Inches(0.30), "RESEARCH GOALS", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(8.5))

    tb_obj = slide.shapes.add_textbox(start_x + Inches(0.5), y_card + Inches(0.85), w_main - Inches(1.0), Inches(3.4))
    tf = tb_obj.text_frame

    objectives = [
        ("Steel Surface Defect Segmentation:", "Develop an effective deep learning framework tailored specifically for small and low-contrast steel surface defects."),
        ("Fine-Grained Feature Preservation:", "Preserve subtle micro-defect details that are conventionally lost during aggressive spatial downsampling and pooling operations."),
        ("Boundary Localization Refinement:", "Improve boundary contour definition for faint, blurred defect edges through dedicated boundary learning."),
        ("Class Imbalance Mitigation:", "Address severe foreground-background imbalance where defect-free steel occupies over 98% of total surface area."),
        ("Comprehensive Benchmark Evaluation:", "Rigorously validate the proposed MSBA-Net architecture against existing state-of-the-art segmentation models on the NEU benchmark.")
    ]

    for title, desc in objectives:
        add_clean_paragraph(tf, f"•  {title}", font_size=Pt(13.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf, f"    {desc}", font_size=Pt(11.5), bold=False, color=COLOR_SLATE, space_after=Pt(10))

    # Bottom takeaway card
    y_focus = y_card + Inches(4.35)
    h_focus = Inches(0.65)
    add_card(slide, start_x + Inches(0.4), y_focus, w_main - Inches(0.8), h_focus, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, border_width=Pt(1.2))
    tb_fc = slide.shapes.add_textbox(start_x + Inches(0.6), y_focus + Inches(0.12), w_main - Inches(1.2), Inches(0.4))
    tf_fc = tb_fc.text_frame
    add_clean_paragraph(tf_fc, "Focus: Accurate segmentation of small defects without sacrificing boundary precision.", font_size=Pt(11.5), bold=True, color=COLOR_GREEN, alignment=PP_ALIGN.CENTER, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 2 — OBJECTIVE (50–60 seconds)
Speaker Script:
"Moving to the core objective of our research:

Our primary aim is to develop a specialized deep learning segmentation framework that directly solves the two most stubborn limitations in industrial optical inspection: microscopic defect scale and near-zero boundary contrast.

Specifically, the research pursues five technical objectives:
First, to engineer an effective segmentation network tailored for steel surface defects.
Second, to preserve fine-grained structural features that standard downsampling operations typically discard.
Third, to improve edge localization for weak and blurry defect margins.
Fourth, to overcome the severe class imbalance where over 98% of captured pixels are normal background steel.
And fifth, to rigorously evaluate the architecture on the standardized NEU surface defect database against published peer-reviewed models.

In short, our central focus is: accurate segmentation of small defects without sacrificing boundary precision."
""")


# =============================================================
# SLIDE 3: BRIEF OVERVIEW / BACKGROUND
# =============================================================
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Brief Overview", "Industrial context, production constraints, and visual segmentation challenges.")

    start_x = Inches(0.8)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Container: Bullet Points (Width: 7.2")
    w_left = Inches(7.2)
    add_card(slide, start_x, y_card, w_left, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.4), Inches(0.30), "INDUSTRIAL CONTEXT", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.5))

    tb_bg = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.85), w_left - Inches(0.8), Inches(4.1))
    tf_bg = tb_bg.text_frame

    bullets = [
        ("Automated Visual Inspection:", "Automated visual inspection is vital for detecting surface defects in modern steel manufacturing, where continuous strip production operates at high speeds."),
        ("Micro-Scale Flaws (<1% Area):", "Steel defects can be extremely small, frequently occupying less than 1% of the entire image canvas (e.g., hairline crazing cracks or tiny pit clusters)."),
        ("Low Intensity Contrast:", "Weak intensity contrast between defect regions and surrounding rolled steel texture makes traditional thresholding and generic segmentation models ineffective."),
        ("Downsampling Information Loss:", "Conventional encoder-decoder networks lose fine spatial information during successive pooling and strided convolutions, causing micro-defects to vanish."),
        ("Need for Pixel Precision:", "Accurate pixel-level segmentation is essential for downstream defect classification, severity assessment, and preventing catastrophic mechanical failures.")
    ]

    for title, desc in bullets:
        add_clean_paragraph(tf_bg, f"•  {title}", font_size=Pt(12.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_bg, f"    {desc}", font_size=Pt(10.0), bold=False, color=COLOR_SLATE, space_after=Pt(8))

    # Right Container: Small Supporting Defect Image (Width: 4.18")
    x_right = start_x + w_left + Inches(0.35)
    w_right = Inches(4.18)
    add_card(slide, x_right, y_card, w_right, h_card, bg_color=COLOR_CARD_ALT, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_right + Inches(0.3), y_card + Inches(0.35), Inches(2.4), Inches(0.30), "DEFECT VISUAL EVIDENCE", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(8.5))

    if os.path.exists(SLIDE3_DEFECT_PATH):
        slide.shapes.add_picture(SLIDE3_DEFECT_PATH, x_right + Inches(0.3), y_card + Inches(0.85), w_right - Inches(0.6), Inches(1.8))

    tb_cap = slide.shapes.add_textbox(x_right + Inches(0.3), y_card + Inches(2.8), w_right - Inches(0.6), Inches(2.1))
    tf_cap = tb_cap.text_frame
    add_clean_paragraph(tf_cap, "Representative Defect Challenge", font_size=Pt(11.0), bold=True, color=COLOR_NAVY, space_after=Pt(3))
    add_clean_paragraph(tf_cap, "Above: Crazing micro-crack from the NEU dataset.", font_size=Pt(9.0), bold=True, color=COLOR_AMBER, space_after=Pt(3))
    add_clean_paragraph(tf_cap, "Notice how the defect blends into the ambient rolled steel texture. The intensity contrast between the crack and background is subtle, and the crack width is merely 1 to 2 pixels.", font_size=Pt(9.0), bold=False, color=COLOR_SLATE, space_after=Pt(4))
    add_clean_paragraph(tf_cap, "Standard CNN downsampling decimates these 1–2 pixel features, causing conventional models to output broken or empty masks.", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 3 — BRIEF OVERVIEW / BACKGROUND (55–65 seconds)
Speaker Script:
"Now, let us examine the industrial background and operational context.

In modern steel manufacturing, flat steel strips are produced in continuous high-speed mills. At these production rates, manual human inspection is physically impossible, making automated optical inspection mandatory.

However, steel surface inspection poses severe computer vision challenges:
First, micro-scale defects such as hairline crazing cracks often occupy less than 1% of the entire image.
Second, defects exhibit remarkably low intensity contrast against the textured steel background. The transition between defect and metal is often just a gradual gradient of a few gray levels.
Third, standard deep learning models rely on aggressive downsampling to expand receptive fields. In doing so, they inadvertently obliterate fine spatial information—a 2-pixel crack completely disappears after a few pooling layers.

On the right side of the slide, you see an actual crazing defect from our research paper. The crack is subtle, hairline-thin, and blends into the surrounding grain. This visual illustrates why standard segmentation models fail, and why pixel-level precision is indispensable."
""")


# =============================================================
# SLIDE 4: PROBLEM STATEMENT & RESEARCH GAP
# =============================================================
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Problem Statement & Research Gap", "Diagnostic analysis of existing limitations and identified architectural needs.")

    start_x = Inches(0.8)
    w_col = Inches(5.69)
    gap_col = Inches(0.35)
    y_card = Inches(1.65)
    h_card = Inches(4.55)

    # Left Column: Key Challenges
    add_card(slide, start_x, y_card, w_col, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.2), Inches(0.30), "KEY CHALLENGES", text_color=COLOR_RED, bg_color=COLOR_RED_BG, border_color=COLOR_RED, font_size=Pt(8.5))

    tb_l = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.85), w_col - Inches(0.8), Inches(3.4))
    tf_l = tb_l.text_frame

    chal = [
        ("Small / Micro-Scale Defects:", "Hairline cracks and tiny pit cavities disappear during repeated downsampling layers."),
        ("Low-Contrast Boundaries:", "Weak intensity transitions make defect edges difficult to distinguish from steel texture."),
        ("Extreme Class Imbalance:", "Over 98% of pixels belong to defect-free steel, biasing standard gradient optimization."),
        ("Feature Dilution:", "Standard CNNs expand receptive fields by pooling, diluting localized high-frequency spatial cues."),
        ("Boundary Bleeding:", "Coarse segmentation masks leak into background steel, degrading boundary accuracy.")
    ]
    for t, d in chal:
        add_clean_paragraph(tf_l, f"•  {t}", font_size=Pt(11.5), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_l, f"    {d}", font_size=Pt(9.5), bold=False, color=COLOR_SLATE, space_after=Pt(6))

    # Right Column: Research Gap
    x_r = start_x + w_col + gap_col
    add_card(slide, x_r, y_card, w_col, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.4), y_card + Inches(0.35), Inches(2.2), Inches(0.30), "RESEARCH GAP", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(8.5))

    tb_r = slide.shapes.add_textbox(x_r + Inches(0.4), y_card + Inches(0.85), w_col - Inches(0.8), Inches(3.4))
    tf_r = tb_r.text_frame

    gaps = [
        ("Feature Loss in Deep Layers:", "Existing segmentation models prioritize macro-context and lose fine structural details."),
        ("Absence of Boundary Guidance:", "Boundary information is rarely preserved explicitly or used to guide decoder upsampling."),
        ("Sub-Optimal Tiny-Defect Performance:", "Prior models achieve high overall mIoU but suffer a steep drop on sub-1% defect areas."),
        ("Lack of Lightweight Synergy:", "A compact approach is needed that jointly preserves multi-scale features and refines boundaries."),
        ("Limited Real-Time Feasibility:", "Complex transformer models often introduce high latency incompatible with factory speeds.")
    ]
    for t, d in gaps:
        add_clean_paragraph(tf_r, f"•  {t}", font_size=Pt(11.5), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_r, f"    {d}", font_size=Pt(9.5), bold=False, color=COLOR_SLATE, space_after=Pt(6))

    # Bottom Takeaway Card
    y_bot = y_card + h_card + Inches(0.12)
    h_bot = Inches(0.58)
    w_full = Inches(11.73)
    add_card(slide, start_x, y_bot, w_full, h_bot, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, border_width=Pt(1.2))
    tb_b = slide.shapes.add_textbox(start_x + Inches(0.4), y_bot + Inches(0.10), w_full - Inches(0.8), Inches(0.38))
    tf_b = tb_b.text_frame
    add_clean_paragraph(tf_b, "Research Need: Multi-scale feature preservation + explicit boundary guidance in a lightweight architecture.", font_size=Pt(11.0), bold=True, color=COLOR_BLUE, alignment=PP_ALIGN.CENTER, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 4 — PROBLEM STATEMENT & RESEARCH GAP (55–65 seconds)
Speaker Script:
"Now that we understand the industrial context, Slide 4 outlines the problem statement and research gap.

On the left, we summarize the key technical challenges:
Micro-scale defects disappear during downsampling; low contrast produces ambiguous edges; more than 98% background pixels create extreme class imbalance; and standard pooling causes feature dilution and boundary bleeding.

On the right, we identify the research gap in existing literature:
First, existing deep architectures like standard U-Net or DeepLab focus primarily on semantic classification and lose fine spatial resolution.
Second, boundary information is rarely preserved explicitly or fed back into the decoder.
Third, while existing models report acceptable overall accuracy on large defects, their performance collapses on tiny defects occupying less than 1% of the image.
And fourth, heavy transformer models introduce latency overheads that exceed real-time line speed limits.

This establishes our clear research need: We require a lightweight architecture that jointly provides multi-scale feature preservation and explicit boundary guidance."
""")


# =============================================================
# SLIDE 5: PROPOSED METHODOLOGY — MSBA-NET (THE ONLY DIAGRAM SLIDE)
# =============================================================
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Proposed Methodology — MSBA-Net", "Clean architectural pipeline: Multi-scale feature preservation and boundary-aware learning.")

    start_x = Inches(0.8)
    w_full = Inches(11.73)
    y_card = Inches(1.65)
    h_card = Inches(5.25)
    add_card(slide, start_x, y_card, w_full, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))

    add_pill_badge(slide, start_x + Inches(0.3), y_card + Inches(0.2), Inches(3.2), Inches(0.28), "ARCHITECTURE SCHEMATIC (FIGURE 1)", text_color=COLOR_NAVY, bg_color=COLOR_CARD_ALT, border_color=COLOR_BORDER, font_size=Pt(8.5))

    # --- ARCHITECTURE FLOWCHART (CLEAN, EDITABLE SHAPES) ---
    # Top Auxiliary Stream: Boundary-Aware Stream (BAS)
    y_bas = y_card + Inches(0.65)
    h_box = Inches(0.72)

    add_card(slide, Inches(4.2), y_bas, Inches(1.8), h_box, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(4.25), y_bas + Inches(0.08), Inches(1.7), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "BAS Stream", font_size=Pt(9.0), bold=True, color=COLOR_AMBER, space_after=Pt(1))
    add_clean_paragraph(tf, "Early Edge Extraction\n(Up-Convolutions)", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    arr_b1 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.1), y_bas + Inches(0.26), Inches(0.3), Inches(0.18))
    arr_b1.fill.solid()
    arr_b1.fill.fore_color.rgb = COLOR_AMBER
    arr_b1.line.fill.background()

    add_card(slide, Inches(6.5), y_bas, Inches(1.8), h_box, bg_color=COLOR_RED_BG, border_color=COLOR_RED, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(6.55), y_bas + Inches(0.08), Inches(1.7), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Boundary Head", font_size=Pt(9.0), bold=True, color=COLOR_RED, space_after=Pt(1))
    add_clean_paragraph(tf, "Contour Prediction\nP_bound (256x256)", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    arr_b2 = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.4), y_bas + Inches(0.26), Inches(0.4), Inches(0.18))
    arr_b2.fill.solid()
    arr_b2.fill.fore_color.rgb = COLOR_PURPLE
    arr_b2.line.fill.background()

    add_card(slide, Inches(8.9), y_bas, Inches(1.8), h_box, bg_color=COLOR_PURPLE_BG, border_color=COLOR_PURPLE, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(8.95), y_bas + Inches(0.08), Inches(1.7), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "BGFF Attention", font_size=Pt(9.0), bold=True, color=COLOR_PURPLE, space_after=Pt(1))
    add_clean_paragraph(tf, "Spatial Attention Gate\n(Residual Guidance)", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    # Middle Main Stream: Backbone Encoder-Decoder
    y_mid = y_card + Inches(1.65)

    add_card(slide, Inches(1.1), y_mid, Inches(1.3), h_box, bg_color=COLOR_CARD_ALT, border_color=COLOR_BORDER, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(1.15), y_mid + Inches(0.08), Inches(1.2), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Input Image", font_size=Pt(9.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
    add_clean_paragraph(tf, "Grayscale Steel\n256 × 256", font_size=Pt(7.5), bold=False, color=COLOR_SLATE, space_after=Pt(0))

    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(2.45), y_mid + Inches(0.26), Inches(0.25), Inches(0.18))
    arr.fill.solid()
    arr.fill.fore_color.rgb = COLOR_BLUE
    arr.line.fill.background()

    add_card(slide, Inches(2.75), y_mid, Inches(1.6), h_box, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(2.80), y_mid + Inches(0.08), Inches(1.5), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Encoder Stage 1–2", font_size=Pt(9.0), bold=True, color=COLOR_BLUE, space_after=Pt(1))
    add_clean_paragraph(tf, "+ MSFP Blocks\nHigh Resolution", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.4), y_mid + Inches(0.26), Inches(0.25), Inches(0.18))
    arr.fill.solid()
    arr.fill.fore_color.rgb = COLOR_BLUE
    arr.line.fill.background()

    add_card(slide, Inches(4.7), y_mid, Inches(1.6), h_box, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(4.75), y_mid + Inches(0.08), Inches(1.5), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Encoder Stage 3–4", font_size=Pt(9.0), bold=True, color=COLOR_BLUE, space_after=Pt(1))
    add_clean_paragraph(tf, "+ MSFP Blocks\nDeep Context", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(6.35), y_mid + Inches(0.26), Inches(0.25), Inches(0.18))
    arr.fill.solid()
    arr.fill.fore_color.rgb = COLOR_NAVY
    arr.line.fill.background()

    add_card(slide, Inches(6.65), y_mid, Inches(1.4), h_box, bg_color=COLOR_CARD_ALT, border_color=COLOR_NAVY, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(6.70), y_mid + Inches(0.08), Inches(1.3), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Bottleneck", font_size=Pt(9.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
    add_clean_paragraph(tf, "Dilated Context\n32 × 32", font_size=Pt(7.5), bold=False, color=COLOR_SLATE, space_after=Pt(0))

    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.1), y_mid + Inches(0.26), Inches(0.25), Inches(0.18))
    arr.fill.solid()
    arr.fill.fore_color.rgb = COLOR_AMBER
    arr.line.fill.background()

    add_card(slide, Inches(8.4), y_mid, Inches(1.5), h_box, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, border_width=Pt(1.2))
    tb = slide.shapes.add_textbox(Inches(8.45), y_mid + Inches(0.08), Inches(1.4), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Decoder", font_size=Pt(9.0), bold=True, color=COLOR_AMBER, space_after=Pt(1))
    add_clean_paragraph(tf, "Guided Up-Conv\n+ Skip Features", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    arr = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(9.95), y_mid + Inches(0.26), Inches(0.25), Inches(0.18))
    arr.fill.solid()
    arr.fill.fore_color.rgb = COLOR_GREEN
    arr.line.fill.background()

    add_card(slide, Inches(10.25), y_mid, Inches(1.7), h_box, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, border_width=Pt(1.4))
    tb = slide.shapes.add_textbox(Inches(10.30), y_mid + Inches(0.08), Inches(1.6), h_box - Inches(0.16))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Defect Mask", font_size=Pt(9.0), bold=True, color=COLOR_GREEN, space_after=Pt(1))
    add_clean_paragraph(tf, "Final Segmentation\n256 × 256", font_size=Pt(7.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    # Connecting vertical arrows for BAS and BGFF
    arr_up = slide.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(3.55), y_mid - Inches(0.35), Inches(0.18), Inches(0.32))
    arr_up.fill.solid()
    arr_up.fill.fore_color.rgb = COLOR_AMBER
    arr_up.line.fill.background()

    arr_down = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.1), y_bas + h_box + Inches(0.05), Inches(0.18), Inches(0.32))
    arr_down.fill.solid()
    arr_down.fill.fore_color.rgb = COLOR_PURPLE
    arr_down.line.fill.background()

    # --- EXPLANATORY BULLETS BELOW DIAGRAM ---
    y_bullets = y_card + Inches(2.65)
    tb_b = slide.shapes.add_textbox(start_x + Inches(0.4), y_bullets, w_full - Inches(0.8), Inches(2.4))
    tf_b = tb_b.text_frame

    arch_bullets = [
        ("MSFP (Multi-Scale Feature Preservation):", "Replaces standard convolution layers in the encoder with parallel dilated depthwise convolutions (dilation rates r = 1, 2, 4) to capture micro-cracks and broad patches simultaneously without gridding artifacts."),
        ("BAS (Boundary-Aware Stream):", "Taps shallow high-resolution encoder features (Stages 1 and 2) to predict an explicit 1-pixel morphological edge map (B_GT), forcing the network to preserve high-frequency contour gradients."),
        ("BGFF (Boundary-Guided Feature Fusion):", "Computes dynamic spatial attention maps (A_spat = σ(Conv([F_bound; F_dec]))) that inject boundary certainty into decoder features, sharpening defect perimeters and suppressing background noise."),
        ("Composite Loss Optimization:", "Jointly supervises regional completeness and edge accuracy using a balanced objective: L_total = 0.3 L_BCE + 0.3 L_Dice + 0.2 L_Focal + 0.2 L_bound, counteracting severe class imbalance.")
    ]

    for title, desc in arch_bullets:
        add_clean_paragraph(tf_b, f"•  {title}", font_size=Pt(11.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_b, f"    {desc}", font_size=Pt(9.0), bold=False, color=COLOR_SLATE, space_after=Pt(5))

    set_speaker_note(slide, """SLIDE 5 — PROPOSED METHODOLOGY / MSBA-NET (75–90 seconds)
Speaker Script:
"This slide presents the core architectural contribution of our paper: MSBA-Net.

Let us trace the information flow:
In the middle row, the input grayscale steel image passes through a four-stage hierarchical encoder. Rather than using conventional convolutions that lose fine spatial resolution, each encoder stage incorporates our Multi-Scale Feature Preservation (MSFP) block. MSFP uses parallel dilated depthwise convolutions with dilation rates 1, 2, and 4. Branch 2 with r=1 has a 3x3 receptive field for hairline micro-cracks; Branch 3 with r=2 spans 5x5 pixels for pit depressions; and Branch 4 with r=4 covers 9x9 pixels for broad patches.

Now look at the top auxiliary pathway: the Boundary-Aware Stream (BAS). BAS branches off early from Stages 1 and 2, where fine edge gradients are still intact. It predicts an explicit 256x256 boundary map under morphological edge supervision.

This learned boundary knowledge is transferred to the decoder via the BGFF—Boundary-Guided Feature Fusion—spatial attention gate. BGFF computes dynamic spatial attention weights that amplify decoder features along uncertain defect boundaries while suppressing noise from uniform steel background.

Finally, the entire network is trained end-to-end using our composite loss: 0.3 BCE, 0.3 Dice, 0.2 Focal, and 0.2 Boundary loss. This joint optimization guarantees both region completeness and razor-sharp boundary delineation."
""")


# =============================================================
# SLIDE 6: EXPERIMENTAL WORK: DATASET & SETUP
# =============================================================
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Experimental Work — Dataset & Setup", "Standardized metallurgical benchmark, data distribution, and defect taxonomy.")

    start_x = Inches(0.8)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Container: Dataset Overview (Width: 5.6")
    w_left = Inches(5.6)
    add_card(slide, start_x, y_card, w_left, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "DATASET SPECIFICATIONS", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.5))

    tb_l = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.85), w_left - Inches(0.8), Inches(4.1))
    tf_l = tb_l.text_frame

    ds_bullets = [
        ("Benchmark Dataset:", "NEU Surface Defect Database developed by Northeastern University, the gold standard in metallurgical computer vision."),
        ("Sample Volume:", "1,800 grayscale industrial images spanning 6 typical defect categories, with exactly 300 images per category."),
        ("Image Dimensions:", "Original resolution 200 × 200 pixels, standardized to 256 × 256 via bilinear interpolation."),
        ("Data Partitioning:", "Strict 70% Training (1,260 images), 15% Validation (270 images), and 15% Testing (270 images)."),
        ("Reproducibility:", "Partitioned using fixed random seeds to guarantee completely held-out, unbiased evaluation.")
    ]

    for title, desc in ds_bullets:
        add_clean_paragraph(tf_l, f"•  {title}", font_size=Pt(11.5), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_l, f"    {desc}", font_size=Pt(9.5), bold=False, color=COLOR_SLATE, space_after=Pt(7))

    # Right Container: 6 Defect Categories Table (Width: 5.78")
    x_r = start_x + w_left + Inches(0.35)
    w_r = Inches(5.78)
    add_card(slide, x_r, y_card, w_r, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "SIX DEFECT CATEGORIES", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(8.5))

    tb_th = slide.shapes.add_textbox(x_r + Inches(0.3), y_card + Inches(0.75), w_r - Inches(0.6), Inches(0.3))
    tf_th = tb_th.text_frame
    add_clean_paragraph(tf_th, "NEU DEFECT TAXONOMY & MORPHOLOGICAL CHALLENGES", font_size=Pt(9.5), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Clean Table
    rows = 7
    cols = 3
    table_shape = slide.shapes.add_table(rows, cols, x_r + Inches(0.3), y_card + Inches(1.10), w_r - Inches(0.6), Inches(3.85))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.50)
    tbl.columns[1].width = Inches(1.68)
    tbl.columns[2].width = Inches(2.00)

    headers = ["Category", "Description", "Primary Challenge"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_WHITE

    cat_data = [
        ("Crazing (Cr)", "Network of micro-cracks", "Hairline cracks (<1% area) vanish"),
        ("Inclusion (In)", "Embedded non-metallic flaws", "Weak contrast & elongated edges"),
        ("Patches (Pa)", "Plate surface peeling", "Irregular boundaries merge with steel"),
        ("Pitted Surface (PS)", "Oxidation cavity clusters", "Dense clusters of tiny depressions"),
        ("Rolled-in Scale (RS)", "Pressed-in oxide sheets", "Fish-scale sheets with irregular contours"),
        ("Scratches (Sc)", "Longitudinal abrasions", "High-aspect-ratio continuous lines")
    ]

    for i, row in enumerate(cat_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_ALT if i % 2 == 1 else COLOR_CARD_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_PRIMARY
            p.font.size = Pt(8.0)
            p.font.bold = (j == 0)
            p.font.color.rgb = COLOR_NAVY if j == 0 else COLOR_SLATE

    set_speaker_note(slide, """SLIDE 6 — EXPERIMENTAL WORK: DATASET & SETUP (60–70 seconds)
Speaker Script:
"Moving on to our experimental setup:

To evaluate MSBA-Net, we benchmarked the architecture on the Northeastern University (NEU) Surface Defect Database, which is the standard benchmark in metallurgical computer vision.

As outlined on the left, the database contains 1,800 grayscale images spanning six defect categories, with 300 images per class. The original 200 by 200 pixel images were resized to 256 by 256. We applied a standard 70-15-15 split: 1,260 images for training, 270 for validation, and 270 for completely held-out test evaluation, using fixed random seeds for strict reproducibility.

On the right, the table outlines the six defect categories: Crazing, Inclusions, Patches, Pitted Surfaces, Rolled-in Scale, and Scratches. Each category introduces distinct morphological challenges: Crazing features micro-cracks that vanish during downsampling; Inclusions exhibit faint contrast gradients; and Pitted Surfaces consist of dense clusters of tiny sub-1% depressions. This diversity ensures a rigorous and comprehensive evaluation."
""")


# =============================================================
# SLIDE 7: EXPERIMENTAL WORK: TRAINING & EVALUATION
# =============================================================
def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Experimental Work — Training & Evaluation", "Hardware configuration, optimization parameters, and quantitative evaluation metrics.")

    start_x = Inches(0.8)
    w_col = Inches(5.69)
    gap_col = Inches(0.35)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Column: Training Configuration & Hardware Setup
    add_card(slide, start_x, y_card, w_col, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "TRAINING & HARDWARE", text_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, font_size=Pt(8.5))

    tb_tr = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.85), w_col - Inches(0.8), Inches(4.1))
    tf_tr = tb_tr.text_frame

    train_bullets = [
        ("Framework & Hardware:", "Implemented in PyTorch 2.1; trained on an enterprise workstation with an NVIDIA RTX 4090 GPU (24GB VRAM)."),
        ("Optimizer & Learning Rate:", "AdamW optimizer (β1=0.9, β2=0.999, weight decay 1e-4) with initial learning rate 1×10⁻³."),
        ("Learning Rate Schedule:", "Cosine Annealing schedule decaying smoothly down to η_min = 1×10⁻⁶."),
        ("Epochs & Batch Size:", "100 training epochs with mini-batch size 16 (reliable convergence observed around epoch 75)."),
        ("Model Parameters:", "9.38 Million parameters (only +1.62M over baseline U-Net backbone)."),
        ("Computational FLOPs:", "14.2 GFLOPs at 256 × 256 input resolution."),
        ("Real-Time Inference Speed:", "41.1 FPS in PyTorch FP32; 88.5 FPS under TensorRT FP16 quantization (easily exceeding 30 FPS line requirements).")
    ]

    for title, desc in train_bullets:
        add_clean_paragraph(tf_tr, f"•  {title}", font_size=Pt(11.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_tr, f"    {desc}", font_size=Pt(9.0), bold=False, color=COLOR_SLATE, space_after=Pt(5))

    # Right Column: Evaluation Metrics
    x_r = start_x + w_col + gap_col
    add_card(slide, x_r, y_card, w_col, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "EVALUATION METRICS", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.5))

    tb_ev = slide.shapes.add_textbox(x_r + Inches(0.4), y_card + Inches(0.85), w_col - Inches(0.8), Inches(4.1))
    tf_ev = tb_ev.text_frame

    metrics = [
        ("Mean Intersection-over-Union (mIoU):", "Primary area-overlap metric measuring segmentation completeness: mIoU = (1/N) Σ [|P ∩ Y| / |P ∪ Y|]."),
        ("Dice Similarity Coefficient:", "Measures harmonic mean of precision and recall; evaluates spatial region overlap and mask continuity."),
        ("Precision & Recall:", "Precision quantifies false positive suppression (97.57%); Recall evaluates defect detection sensitivity (99.41%)."),
        ("Tiny-Defect mIoU (Critical Metric):", "Explicitly isolates performance on micro-defects occupying less than 1% of total image area (<655 pixels).")
    ]

    for title, desc in metrics:
        add_clean_paragraph(tf_ev, f"•  {title}", font_size=Pt(11.5), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_ev, f"    {desc}", font_size=Pt(9.5), bold=False, color=COLOR_SLATE, space_after=Pt(8))

    # Tiny Defect Definition Callout
    y_call = y_card + Inches(3.85)
    add_card(slide, x_r + Inches(0.4), y_call, w_col - Inches(0.8), Inches(1.05), bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, border_width=Pt(1.2))
    tb_c = slide.shapes.add_textbox(x_r + Inches(0.6), y_call + Inches(0.12), w_col - Inches(1.2), Inches(0.8))
    tf_c = tb_c.text_frame
    add_clean_paragraph(tf_c, "Why Tiny-Defect mIoU is Critical:", font_size=Pt(10.0), bold=True, color=COLOR_AMBER, space_after=Pt(2))
    add_clean_paragraph(tf_c, "Defects occupying <1% area are the primary cause of fatigue tears in automotive stamping. A model can score high on large defects while completely missing micro-cracks. Tiny mIoU unmasks this failure mode.", font_size=Pt(8.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 7 — EXPERIMENTAL WORK: TRAINING & EVALUATION (60–70 seconds)
Speaker Script:
"Now, let us examine the training setup and evaluation metrics shown on Slide 7.

On the left, all models were trained in PyTorch 2.1 on an NVIDIA RTX 4090 GPU with 24GB of VRAM. We used the AdamW optimizer with an initial learning rate of 1 times 10 to the minus 3, decaying via cosine annealing down to 1 times 10 to the minus 6, over 100 epochs with a batch size of 16. Convergence was reliably reached by epoch 75.

From an engineering perspective, MSBA-Net maintains a lightweight footprint of only 9.38 million parameters and 14.2 GFLOPs. It achieves 41.1 frames per second in standard FP32 PyTorch, and 88.5 frames per second under TensorRT FP16 quantization—easily satisfying real-time factory line requirements.

On the right, we define our evaluation metrics: mIoU, Dice, Precision, and Recall. Crucially, we introduce Tiny-Defect mIoU, which isolates defect regions occupying less than 1% of the image canvas. In industrial steel production, a tiny hairline crack can initiate catastrophic structural failure during stamping. Generic models often score high overall because large defects mask their poor performance on micro-flaws. Tiny-Defect mIoU directly measures whether micro-scale features are truly preserved."
""")


# =============================================================
# SLIDE 8: RESULTS: COMPARISON WITH EXISTING METHODS
# =============================================================
def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Results — Comparison with Existing Methods", "Comparative benchmark evaluation against peer-reviewed 2024–2026 segmentation models.")

    start_x = Inches(0.8)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Container: Editable Column Chart for mIoU (Width: 6.2")
    w_chart = Inches(6.2)
    add_card(slide, start_x, y_card, w_chart, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.3), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "BENCHMARK COMPARISON", text_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, font_size=Pt(8.5))

    tb_ct = slide.shapes.add_textbox(start_x + Inches(0.3), y_card + Inches(0.75), w_chart - Inches(0.6), Inches(0.3))
    tf_ct = tb_ct.text_frame
    add_clean_paragraph(tf_ct, "mIoU (%) ACROSS PUBLISHED BENCHMARK MODELS", font_size=Pt(10.0), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Native PowerPoint Chart
    chart_data = CategoryChartData()
    chart_data.categories = ['GNA-PSC', 'MixSegNet', 'EF²Net', 'DAMF-Net', 'DMS-UNet', 'MSBA-Net']
    chart_data.add_series('mIoU (%)', (83.50, 84.80, 88.42, 91.65, 94.20, 97.01))

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, start_x + Inches(0.3), y_card + Inches(1.10), w_chart - Inches(0.6), Inches(3.85), chart_data
    )
    chart = chart_shape.chart
    chart.has_legend = False

    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.font.name = FONT_PRIMARY
    data_labels.font.size = Pt(9.0)
    data_labels.font.bold = True
    data_labels.font.color.rgb = COLOR_NAVY
    data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    series = plot.series[0]
    for idx, pt in enumerate(series.points):
        pt.format.fill.solid()
        if idx == 5:
            pt.format.fill.fore_color.rgb = COLOR_GREEN
        elif idx == 4:
            pt.format.fill.fore_color.rgb = COLOR_BLUE
        else:
            pt.format.fill.fore_color.rgb = COLOR_MUTED

    # Right Container: Compact Table & Key Observations (Width: 5.18")
    x_r = start_x + w_chart + Inches(0.35)
    w_r = Inches(5.18)
    add_card(slide, x_r, y_card, w_r, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.3), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "QUANTITATIVE SUMMARY", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.5))

    tb_tt = slide.shapes.add_textbox(x_r + Inches(0.3), y_card + Inches(0.75), w_r - Inches(0.6), Inches(0.3))
    tf_tt = tb_tt.text_frame
    add_clean_paragraph(tf_tt, "PERFORMANCE METRICS BREAKDOWN", font_size=Pt(10.0), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Compact Table
    rows = 7
    cols = 4
    table_shape = slide.shapes.add_table(rows, cols, x_r + Inches(0.3), y_card + Inches(1.10), w_r - Inches(0.6), Inches(2.20))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.68)
    tbl.columns[1].width = Inches(0.95)
    tbl.columns[2].width = Inches(0.95)
    tbl.columns[3].width = Inches(1.00)

    headers = ["Model", "mIoU", "Dice", "Tiny mIoU"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_WHITE
        p.alignment = PP_ALIGN.CENTER

    data = [
        ("GNA-PSC (2024)", "83.50%", "88.20%", "72.40%"),
        ("MixSegNet (2024)", "84.80%", "91.50%", "76.10%"),
        ("EF²Net (2024)", "88.42%", "93.85%", "81.25%"),
        ("DAMF-Net (2025)", "91.65%", "95.60%", "85.70%"),
        ("DMS-UNet (2026)", "94.20%", "97.02%", "88.40%"),
        ("MSBA-Net (Ours)", "97.01%", "98.48%", "92.48%")
    ]

    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_GREEN_BG if i == 5 else (COLOR_CARD_ALT if i % 2 == 1 else COLOR_CARD_WHITE)
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_PRIMARY
            p.font.size = Pt(8.0)
            p.font.bold = (i == 5 or j == 0)
            p.font.color.rgb = COLOR_GREEN if i == 5 else COLOR_NAVY
            if j > 0:
                p.alignment = PP_ALIGN.CENTER

    # Observations below table
    tb_obs = slide.shapes.add_textbox(x_r + Inches(0.3), y_card + Inches(3.45), w_r - Inches(0.6), Inches(1.6))
    tf_obs = tb_obs.text_frame
    add_clean_paragraph(tf_obs, "Key Observations:", font_size=Pt(10.5), bold=True, color=COLOR_NAVY, space_after=Pt(3))
    add_clean_paragraph(tf_obs, "• MSBA-Net achieves 97.01% mIoU, establishing a new SOTA on the NEU benchmark (+2.81% over 2026 DMS-UNet).", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(4))
    add_clean_paragraph(tf_obs, "• Dice coefficient reaches 98.48%, with 97.57% Precision and 99.41% Recall (near-zero missed flaws).", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(4))
    add_clean_paragraph(tf_obs, "• Tiny-Defect mIoU reaches 92.48% (+4.08% over DMS-UNet and +20.08% over GNA-PSC), proving micro-flaw preservation.", font_size=Pt(9.0), bold=False, color=COLOR_GREEN, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 8 — RESULTS: COMPARISON WITH EXISTING METHODS (70–80 seconds)
Speaker Script:
"Now we compare the proposed method with existing approaches on Slide 8.

Looking at the column chart on the left, MSBA-Net establishes a decisive lead across published baselines:
Earlier 2024 models like GNA-PSC and MixSegNet achieved 83.50% and 84.80% mIoU. More recent models like EF²Net in 2024 and DAMF-Net in 2025 reached 88.42% and 91.65%. The previous state-of-the-art, DMS-UNet published in 2026, achieved 94.20%.
MSBA-Net advances the benchmark to 97.01% mIoU, an improvement of plus 2.81 percentage points.

Examining the summary table on the right:
Our Dice coefficient reaches 98.48%, indicating outstanding mask continuity and region overlap.
Furthermore, MSBA-Net achieves 97.57% Precision and 99.41% Recall, translating to near-zero missed defects in practice.
Most significantly, look at the final column: Tiny-Defect mIoU. On defects occupying less than 1% of the image, MSBA-Net achieves 92.48%. That is a 4.08% improvement over DMS-UNet and more than 20% higher than GNA-PSC. This provides clear empirical proof that our hybrid multi-scale and boundary-aware approach effectively prevents micro-scale feature dilution."
""")


# =============================================================
# SLIDE 9: RESULTS: DETAILED ANALYSIS / PER-CLASS PERFORMANCE
# =============================================================
def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Results — Detailed Analysis", "Class-wise performance breakdown across all six defect categories and visual validation.")

    start_x = Inches(0.8)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Container: Per-Class Table & Observations (Width: 5.8")
    w_left = Inches(5.8)
    add_card(slide, start_x, y_card, w_left, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "PER-CLASS ACCURACY", text_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, font_size=Pt(8.5))

    tb_ph = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.75), w_left - Inches(0.8), Inches(0.3))
    tf_ph = tb_ph.text_frame
    add_clean_paragraph(tf_ph, "PER-CATEGORY SEGMENTATION ACCURACY (NEU TEST SET)", font_size=Pt(9.5), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Per-Class Table
    rows = 7
    cols = 3
    table_shape = slide.shapes.add_table(rows, cols, start_x + Inches(0.4), y_card + Inches(1.10), w_left - Inches(0.8), Inches(2.20))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(1.4)
    tbl.columns[2].width = Inches(1.4)

    headers = ["Defect Category", "mIoU (%)", "Dice (%)"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_WHITE
        p.alignment = PP_ALIGN.CENTER

    per_class_data = [
        ("Patches (Pa)", "97.85%", "98.91%"),
        ("Inclusion (In)", "97.15%", "98.55%"),
        ("Scratches (Sc)", "97.12%", "98.54%"),
        ("Pitted Surface (PS)", "96.80%", "98.37%"),
        ("Rolled-in Scale (RS)", "96.72%", "98.33%"),
        ("Crazing (Cr)", "96.42%", "98.18%")
    ]

    for i, row in enumerate(per_class_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_CARD_ALT if i % 2 == 1 else COLOR_CARD_WHITE
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_PRIMARY
            p.font.size = Pt(8.0)
            p.font.bold = (j == 0)
            p.font.color.rgb = COLOR_NAVY
            if j > 0:
                p.alignment = PP_ALIGN.CENTER

    # Observations below table
    tb_co = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(3.45), w_left - Inches(0.8), Inches(1.6))
    tf_co = tb_co.text_frame
    add_clean_paragraph(tf_co, "Key Observations:", font_size=Pt(10.5), bold=True, color=COLOR_NAVY, space_after=Pt(3))
    add_clean_paragraph(tf_co, "• Consistent High Performance: Every defect class achieves over 96.4% mIoU, demonstrating robust generalization across diverse defect morphologies.", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(4))
    add_clean_paragraph(tf_co, "• Patches & Inclusions: Achieve peak accuracy (97.85% and 97.15%) due to continuous regional texture and distinct contrast boundaries.", font_size=Pt(9.0), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(4))
    add_clean_paragraph(tf_co, "• Crazing Micro-Cracks: Comparatively challenging (96.42%) due to extreme hairline thinness, but still exceeds prior benchmarks by over 3%.", font_size=Pt(9.0), bold=False, color=COLOR_AMBER, space_after=Pt(0))

    # Right Container: Small Supporting Qualitative Comparison (Width: 5.58")
    x_r = start_x + w_left + Inches(0.35)
    w_r = Inches(5.58)
    add_card(slide, x_r, y_card, w_r, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "QUALITATIVE EVIDENCE", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.5))

    if os.path.exists(SLIDE9_QUAL_PATH):
        slide.shapes.add_picture(SLIDE9_QUAL_PATH, x_r + Inches(0.3), y_card + Inches(0.85), w_r - Inches(0.6), Inches(2.6))

    tb_qc = slide.shapes.add_textbox(x_r + Inches(0.3), y_card + Inches(3.60), w_r - Inches(0.6), Inches(1.4))
    tf_qc = tb_qc.text_frame
    add_clean_paragraph(tf_qc, "Visual Progression Insights (Figure 2 Crop):", font_size=Pt(10.0), bold=True, color=COLOR_NAVY, space_after=Pt(2))
    add_clean_paragraph(tf_qc, "• Row 1 (Crazing): Baseline U-Net produces a fragmented, broken mask (82.9% IoU). MSBA-Net traces the continuous micro-crack trajectory (86.7% IoU).", font_size=Pt(8.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(3))
    add_clean_paragraph(tf_qc, "• Row 2 (Inclusion): Baseline exhibits significant border bleeding into background steel; MSBA-Net matches ground-truth boundaries cleanly.", font_size=Pt(8.5), bold=False, color=COLOR_GREEN, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 9 — RESULTS: DETAILED ANALYSIS (65–75 seconds)
Speaker Script:
"Moving to a more detailed analysis on Slide 9, we examine per-category performance across all six defect classes.

On the left, the table demonstrates remarkable consistency:
Every single defect category surpasses 96.4% mIoU. Patches achieved the highest accuracy at 97.85% mIoU, followed by Inclusions at 97.15% and Scratches at 97.12%.
Crazing was the most challenging category at 96.42% mIoU, owing to the extreme thinness of hairline micro-cracks. Nevertheless, MSBA-Net outperforms all competing models on crazing by over 3 percentage points.

On the right, we show supporting qualitative visual evidence directly from Figure 2 of our paper:
In the first row, representing crazing: the baseline U-Net produces a fragmented, incomplete mask. In contrast, MSBA-Net traces the entire crack trajectory continuously.
In the second row, showing thin inclusions: baseline U-Net bleeds noticeably into the background steel, whereas MSBA-Net confines its mask strictly to the defect perimeter.
This visual comparison confirms that MSFP preserves fine structural topology, while BAS and BGFF prevent boundary bleeding."
""")


# =============================================================
# SLIDE 10: ABLATION STUDY & ANALYSIS
# =============================================================
def build_slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Ablation Study & Analysis", "Component-wise incremental validation of MSFP, BAS, and BGFF modules.")

    start_x = Inches(0.8)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Container: Editable Grouped Column Chart (Width: 6.2")
    w_chart = Inches(6.2)
    add_card(slide, start_x, y_card, w_chart, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.3), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "ABLATION PROGRESSION", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(8.5))

    tb_ct = slide.shapes.add_textbox(start_x + Inches(0.3), y_card + Inches(0.72), Inches(2.8), Inches(0.3))
    tf_ct = tb_ct.text_frame
    add_clean_paragraph(tf_ct, "INCREMENTAL mIoU vs TINY mIoU", font_size=Pt(9.5), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Clean custom legend badges
    add_pill_badge(slide, start_x + Inches(3.1), y_card + Inches(0.72), Inches(1.35), Inches(0.24), "■ Overall mIoU", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(7.5))
    add_pill_badge(slide, start_x + Inches(4.55), y_card + Inches(0.72), Inches(1.50), Inches(0.24), "■ Tiny-Defect mIoU", text_color=COLOR_AMBER, bg_color=COLOR_AMBER_BG, border_color=COLOR_AMBER, font_size=Pt(7.5))

    # Native Grouped Bar Chart
    chart_data = CategoryChartData()
    chart_data.categories = ['Baseline Backbone', '+ MSFP Blocks', '+ BAS Stream', '+ BGFF (Full)']
    chart_data.add_series('Overall mIoU (%)', (94.10, 95.85, 96.48, 97.01))
    chart_data.add_series('Tiny-Defect mIoU (%)', (84.60, 89.65, 91.20, 92.48))

    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, start_x + Inches(0.3), y_card + Inches(1.08), w_chart - Inches(0.6), Inches(3.90), chart_data
    )
    chart = chart_shape.chart
    chart.has_legend = False

    plot = chart.plots[0]
    plot.has_data_labels = True
    data_labels = plot.data_labels
    data_labels.font.name = FONT_PRIMARY
    data_labels.font.size = Pt(8.0)
    data_labels.font.bold = True
    data_labels.position = XL_LABEL_POSITION.OUTSIDE_END

    # Colors: Series 0 (Blue), Series 1 (Amber)
    series1 = plot.series[0]
    for pt in series1.points:
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = COLOR_BLUE

    series2 = plot.series[1]
    for pt in series2.points:
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = COLOR_AMBER

    # Right Container: Compact Table & Key Takeaways (Width: 5.18")
    x_r = start_x + w_chart + Inches(0.35)
    w_r = Inches(5.18)
    add_card(slide, x_r, y_card, w_r, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.3), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "ABLATION FINDINGS", text_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, font_size=Pt(8.5))

    tb_tt = slide.shapes.add_textbox(x_r + Inches(0.3), y_card + Inches(0.75), w_r - Inches(0.6), Inches(0.3))
    tf_tt = tb_tt.text_frame
    add_clean_paragraph(tf_tt, "COMPONENT CONTRIBUTION SUMMARY", font_size=Pt(10.0), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Compact Table
    rows = 5
    cols = 5
    table_shape = slide.shapes.add_table(rows, cols, x_r + Inches(0.3), y_card + Inches(1.10), w_r - Inches(0.6), Inches(1.85))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.78)
    tbl.columns[1].width = Inches(0.70)
    tbl.columns[2].width = Inches(0.70)
    tbl.columns[3].width = Inches(0.75)
    tbl.columns[4].width = Inches(0.65)

    headers = ["Configuration", "mIoU", "Dice", "Tiny", "Params"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(8.0)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_WHITE
        p.alignment = PP_ALIGN.CENTER

    ablation_data = [
        ("Baseline Backbone", "94.10%", "96.90%", "84.60%", "7.76M"),
        ("+ MSFP Blocks", "95.85%", "97.88%", "89.65%", "8.42M"),
        ("+ BAS Stream", "96.48%", "98.21%", "91.20%", "9.15M"),
        ("+ BGFF Gate (Full)", "97.01%", "98.48%", "92.48%", "9.38M")
    ]

    for i, row in enumerate(ablation_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_GREEN_BG if i == 3 else (COLOR_CARD_ALT if i % 2 == 1 else COLOR_CARD_WHITE)
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_PRIMARY
            p.font.size = Pt(8.0)
            p.font.bold = (i == 3 or j == 0)
            p.font.color.rgb = COLOR_GREEN if i == 3 else COLOR_NAVY
            if j > 0:
                p.alignment = PP_ALIGN.CENTER

    # 3 Takeaway Bullets below table
    tb_tk = slide.shapes.add_textbox(x_r + Inches(0.3), y_card + Inches(3.10), w_r - Inches(0.6), Inches(1.95))
    tf_tk = tb_tk.text_frame
    add_clean_paragraph(tf_tk, "Core Ablation Insights:", font_size=Pt(10.5), bold=True, color=COLOR_NAVY, space_after=Pt(3))
    add_clean_paragraph(tf_tk, "1. MSFP Drives Micro-Defect Recovery (+5.05% Tiny mIoU): Integrating MSFP blocks produces the single largest surge in tiny-defect mIoU (84.60% → 89.65%), proving parallel dilated convolutions prevent hairline crack decimation.", font_size=Pt(8.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(4))
    add_clean_paragraph(tf_tk, "2. BAS Refines Boundary Supervision (+0.63% mIoU): Explicit edge supervision forces early encoder features to retain high-frequency contour gradients.", font_size=Pt(8.5), bold=False, color=COLOR_NAVY_LIGHT, space_after=Pt(4))
    add_clean_paragraph(tf_tk, "3. BGFF Delivers Peak Fusion (+0.53% mIoU): Spatial attention gating completes the network, reaching 97.01% mIoU with only 1.62M total parameter addition over baseline.", font_size=Pt(8.5), bold=False, color=COLOR_GREEN, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 10 — ABLATION STUDY & ANALYSIS (70–80 seconds)
Speaker Script:
"Why does MSBA-Net work so effectively? Slide 10 provides empirical ablation evidence validating each component.

Looking at the grouped column chart on the left, we track both overall mIoU in blue and tiny-defect mIoU in amber across four incremental configurations:
Starting from our Baseline Backbone without MSFP or BAS, the model achieves 94.10% overall mIoU, but only 84.60% on tiny defects.

When we insert our Multi-Scale Feature Preservation (MSFP) blocks:
Overall mIoU increases to 95.85%, but tiny-defect mIoU surges by an impressive 5.05 percentage points to 89.65%. This confirms our fundamental hypothesis: parallel dilated depthwise convolutions effectively prevent micro-flaws from being washed out.

Next, when we introduce the Boundary-Aware Stream (BAS):
Overall mIoU rises to 96.48%, and tiny mIoU reaches 91.20%, proving that explicit boundary loss forces the early encoder to maintain high-frequency edge gradients.

Finally, adding the BGFF spatial attention gate completes MSBA-Net, achieving our peak 97.01% mIoU and 92.48% tiny-defect mIoU.
As shown in the table on the right, this entire architectural enhancement adds only 1.62 million parameters over baseline, demonstrating high parameter efficiency."
""")


# =============================================================
# SLIDE 11: LIMITATIONS & ROBUSTNESS
# =============================================================
def build_slide_11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Limitations & Robustness Analysis", "Honest reporting of operational boundaries and simulated factory degradation stress-testing.")

    start_x = Inches(0.8)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Container: Degradation Results Table (Width: 5.4")
    w_left = Inches(5.4)
    add_card(slide, start_x, y_card, w_left, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "STRESS-TEST RESULTS", text_color=COLOR_RED, bg_color=COLOR_RED_BG, border_color=COLOR_RED, font_size=Pt(8.5))

    tb_th = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.75), w_left - Inches(0.8), Inches(0.3))
    tf_th = tb_th.text_frame
    add_clean_paragraph(tf_th, "SIMULATED SHOP-FLOOR DEGRADATION BENCHMARK", font_size=Pt(9.5), bold=True, color=COLOR_NAVY, space_after=Pt(0))

    # Clean Table
    rows = 5
    cols = 3
    table_shape = slide.shapes.add_table(rows, cols, start_x + Inches(0.4), y_card + Inches(1.10), w_left - Inches(0.8), Inches(2.20))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.1)
    tbl.columns[1].width = Inches(1.1)
    tbl.columns[2].width = Inches(1.4)

    headers = ["Operating Condition", "mIoU", "Observed Impact"]
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_NAVY
        p = cell.text_frame.paragraphs[0]
        p.font.name = FONT_PRIMARY
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.color.rgb = COLOR_CARD_WHITE
        p.alignment = PP_ALIGN.CENTER

    deg_data = [
        ("Normal Benchmark", "97.01%", "Clean baseline reference"),
        ("Low Contrast (γ = 1.8)", "39.67%", "Moderate degradation"),
        ("Defocus Blur (σ = 3.0)", "26.29%", "Severe edge attenuation"),
        ("Gaussian Noise (σ = 25)", "2.33%", "Complete breakdown")
    ]

    for i, row in enumerate(deg_data):
        for j, val in enumerate(row):
            cell = tbl.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_RED_BG if i == 3 else (COLOR_AMBER_BG if i in [1, 2] else COLOR_GREEN_BG)
            p = cell.text_frame.paragraphs[0]
            p.font.name = FONT_PRIMARY
            p.font.size = Pt(8.0)
            p.font.bold = (j == 0)
            p.font.color.rgb = COLOR_RED if i == 3 else (COLOR_AMBER if i in [1, 2] else COLOR_GREEN)
            if j == 1:
                p.alignment = PP_ALIGN.CENTER

    tb_sub = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(3.45), w_left - Inches(0.8), Inches(1.6))
    tf_sub = tb_sub.text_frame
    add_clean_paragraph(tf_sub, "Scientific Rigor in Reporting:", font_size=Pt(10.0), bold=True, color=COLOR_NAVY, space_after=Pt(2))
    add_clean_paragraph(tf_sub, "Section V of our paper stress-tested the model against optical distortions common in hot strip mills. We report these steep performance drops transparently rather than claiming unconstrained readiness.", font_size=Pt(8.5), bold=False, color=COLOR_SLATE, space_after=Pt(0))

    # Right Container: Limitations Bullet Points (Width: 5.98")
    x_r = start_x + w_left + Inches(0.35)
    w_r = Inches(5.98)
    add_card(slide, x_r, y_card, w_r, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.4), y_card + Inches(0.35), Inches(2.6), Inches(0.30), "IDENTIFIED LIMITATIONS", text_color=COLOR_RED, bg_color=COLOR_RED_BG, border_color=COLOR_RED, font_size=Pt(8.5))

    tb_lim = slide.shapes.add_textbox(x_r + Inches(0.4), y_card + Inches(0.85), w_r - Inches(0.8), Inches(4.1))
    tf_lim = tb_lim.text_frame

    lim_bullets = [
        ("Performance Degradation Under Optical Stress:", "Segmentation accuracy decreases substantially under severe optical degradation (e.g., vapor fog or mechanical flutter)."),
        ("Gaussian Noise Vulnerability (2.33% mIoU):", "Random thermal sensor noise corrupts the subtle 1–2 pixel gray-level transitions required by morphological boundary convolutions, leading to breakdown."),
        ("Information-Theoretic Root Cause:", "Differential edge operations inherently amplify high-frequency random noise; without pre-filtering, noise edges overpower true defect perimeters."),
        ("Laboratory Scope vs. Rolling Mills:", "The current evaluation is primarily based on the static NEU benchmark and simulated perturbations. Moving strip mills involve dynamic vibrations and lighting shifts."),
        ("Industrial Deployment Prerequisites:", "Real-world continuous mill integration requires active LED strobing shrouds, sealed camera housings, and front-end bilateral denoising filters.")
    ]

    for title, desc in lim_bullets:
        add_clean_paragraph(tf_lim, f"•  {title}", font_size=Pt(11.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_lim, f"    {desc}", font_size=Pt(9.0), bold=False, color=COLOR_SLATE, space_after=Pt(6))

    set_speaker_note(slide, """SLIDE 11 — LIMITATIONS / ROBUSTNESS ANALYSIS (60–75 seconds)
Speaker Script:
"Finally, let us discuss the limitations and operational boundaries on Slide 11.

Academic integrity requires complete transparency about physical operational limits. In Section V of our paper, we stress-tested MSBA-Net under simulated factory degradations:
As shown in the table on the left:
Under low contrast (gamma=1.8), mIoU drops to 39.67%.
Under defocus blur (sigma=3.0), simulating camera vibration, mIoU drops to 26.29%.
Under severe Gaussian noise (sigma=25), performance breaks down to 2.33% mIoU.

Why does this breakdown occur?
This is an information-theoretic limitation rather than an architectural bug:
In steel inspection, defect boundaries are defined by subtle differences of only 5 to 15 gray levels. Adding Gaussian noise with sigma=25 produces random fluctuations of up to 50 gray levels across adjacent pixels. Because our Boundary-Aware Stream computes morphological gradients via differential operations, high-frequency random noise generates thousands of spurious false edges, completely overwhelming the attention gate.

This provides an important engineering takeaway: Deploying deep learning models in real steel mills requires active LED shrouding, sealed camera enclosures, and front-end bilateral denoising filters. We do not claim unconstrained factory readiness without appropriate optical engineering."
""")


# =============================================================
# SLIDE 12: CONCLUSIONS & FUTURE WORK
# =============================================================
def build_slide_12(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    add_header(slide, "Research Seminar", "Conclusions & Future Work", "Summary of core contributions, benchmark findings, and future research directions.")

    start_x = Inches(0.8)
    w_col = Inches(5.69)
    gap_col = Inches(0.35)
    y_card = Inches(1.65)
    h_card = Inches(5.25)

    # Left Column: Conclusions
    add_card(slide, start_x, y_card, w_col, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, start_x + Inches(0.4), y_card + Inches(0.35), Inches(2.2), Inches(0.30), "CONCLUSIONS", text_color=COLOR_GREEN, bg_color=COLOR_GREEN_BG, border_color=COLOR_GREEN, font_size=Pt(8.5))

    tb_c = slide.shapes.add_textbox(start_x + Inches(0.4), y_card + Inches(0.85), w_col - Inches(0.8), Inches(4.1))
    tf_c = tb_c.text_frame

    concl = [
        ("MSBA-Net Architecture:", "Successfully formulated an architecture combining multi-scale feature preservation (MSFP) with boundary-aware learning (BAS + BGFF)."),
        ("Micro-Defect Preservation:", "MSFP blocks effectively prevent hairline cracks and sub-1% pit depressions from vanishing during encoder downsampling."),
        ("Boundary Localization:", "BAS boundary supervision and BGFF spatial attention gating eliminate border bleeding and sharpen defect perimeters."),
        ("State-of-the-Art Benchmark:", "Achieved 97.01% mIoU, 98.48% Dice, and 92.48% Tiny-Defect mIoU on the NEU dataset, outperforming all 2024–2026 baselines."),
        ("Real-Time Efficiency:", "Operates at 41.1 FPS in PyTorch FP32 and 88.5 FPS under TensorRT FP16, validating real-world factory feasibility.")
    ]

    for title, desc in concl:
        add_clean_paragraph(tf_c, f"•  {title}", font_size=Pt(11.0), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_c, f"    {desc}", font_size=Pt(9.0), bold=False, color=COLOR_SLATE, space_after=Pt(6))

    # Right Column: Future Work
    x_r = start_x + w_col + gap_col
    add_card(slide, x_r, y_card, w_col, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))
    add_pill_badge(slide, x_r + Inches(0.4), y_card + Inches(0.35), Inches(2.2), Inches(0.30), "FUTURE WORK", text_color=COLOR_BLUE, bg_color=COLOR_BLUE_BG, border_color=COLOR_BLUE, font_size=Pt(8.5))

    tb_f = slide.shapes.add_textbox(x_r + Inches(0.4), y_card + Inches(0.85), w_col - Inches(0.8), Inches(4.1))
    tf_f = tb_f.text_frame

    fut = [
        ("Temporal Line-Scan Video Attention:", "Extend 2D spatial defect segmentation toward spatio-temporal attention across continuous video frames on moving strip coils."),
        ("Self-Supervised Pre-Training (MAE):", "Pre-train masked autoencoders on massive uncurated industrial steel imagery to improve generalized representation before fine-tuning."),
        ("Embedded Edge Deployment (INT8):", "Investigate 8-bit integer quantization and post-training compression for direct deployment on NVIDIA Jetson and FPGA hardware."),
        ("Adaptive Front-End Denoising:", "Integrate lightweight learnable bilateral denoising modules to maintain boundary stability under severe Gaussian sensor noise.")
    ]

    for title, desc in fut:
        add_clean_paragraph(tf_f, f"•  {title}", font_size=Pt(11.5), bold=True, color=COLOR_NAVY, space_after=Pt(1))
        add_clean_paragraph(tf_f, f"    {desc}", font_size=Pt(9.5), bold=False, color=COLOR_SLATE, space_after=Pt(8))

    set_speaker_note(slide, """SLIDE 12 — CONCLUSIONS & FUTURE WORK (60–70 seconds)
Speaker Script:
"To conclude our presentation, Slide 12 summarizes our core findings and outlines future directions.

In conclusion:
First, MSBA-Net successfully unites multi-scale feature preservation with boundary-aware learning.
Second, MSFP prevents hairline cracks and tiny pit depressions from vanishing during downsampling.
Third, BAS and BGFF establish sharp, leak-free boundary contours around faint defect perimeters.
Fourth, the model sets a new benchmark on the NEU dataset, delivering 97.01% mIoU, 98.48% Dice, and 92.48% Tiny-Defect mIoU.
And fifth, it operates at 41.1 FPS in PyTorch and 88.5 FPS under TensorRT FP16, validating practical industrial viability.

Looking forward, our future roadmap focuses on:
1. Extending the model to spatio-temporal video attention across continuous line-scan strip feeds.
2. Exploring self-supervised pre-training on massive unlabeled industrial steel coils.
3. Implementing INT8 edge quantization for low-power industrial controllers.
And 4. Adding adaptive front-end denoising to harden boundary convolutions against sensor noise."
""")


# =============================================================
# SLIDE 13: THANK YOU
# =============================================================
def build_slide_13(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)

    start_x = Inches(1.5)
    w_main = Inches(10.333)
    y_card = Inches(1.5)
    h_card = Inches(4.5)
    add_card(slide, start_x, y_card, w_main, h_card, bg_color=COLOR_CARD_WHITE, border_color=COLOR_BORDER, border_width=Pt(1.5))

    tb = slide.shapes.add_textbox(start_x + Inches(0.8), y_card + Inches(0.8), w_main - Inches(1.6), Inches(3.0))
    tf = tb.text_frame
    add_clean_paragraph(tf, "Thank You", font_name=FONT_HEADING, font_size=Pt(36), bold=True, color=COLOR_NAVY, alignment=PP_ALIGN.CENTER, space_after=Pt(6))
    add_clean_paragraph(tf, "Questions & Discussion", font_size=Pt(18), bold=True, color=COLOR_AMBER, alignment=PP_ALIGN.CENTER, space_after=Pt(20))

    # Divider line
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, start_x + Inches(3.0), y_card + Inches(2.2), w_main - Inches(6.0), Pt(1.5))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_MUTED
    line.line.fill.background()

    tb_sub = slide.shapes.add_textbox(start_x + Inches(0.8), y_card + Inches(2.45), w_main - Inches(1.6), Inches(1.5))
    tf_sub = tb_sub.text_frame
    add_clean_paragraph(tf_sub, "Amen Parmar", font_size=Pt(14), bold=True, color=COLOR_NAVY, alignment=PP_ALIGN.CENTER, space_after=Pt(3))
    add_clean_paragraph(tf_sub, "M.Tech (AIML)  •  Roll No: 2502010010", font_size=Pt(11), bold=False, color=COLOR_SLATE, alignment=PP_ALIGN.CENTER, space_after=Pt(3))
    add_clean_paragraph(tf_sub, "Department of Computer Science and Engineering  |  Manipal University Jaipur", font_size=Pt(10), bold=False, color=COLOR_MUTED, alignment=PP_ALIGN.CENTER, space_after=Pt(0))

    set_speaker_note(slide, """SLIDE 13 — THANK YOU (30 seconds)
Speaker Script:
"Thank you very much, respected committee members and professors, for your time, patience, and kind attention.

I will now be honored to take any questions, comments, or suggestions regarding our research."
""")


# =============================================================
# MAIN FUNCTION
# =============================================================
def main():
    prs = pptx.Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    print("Building Slide 1: Title...")
    build_slide_1(prs)
    print("Building Slide 2: Objective...")
    build_slide_2(prs)
    print("Building Slide 3: Brief Overview / Background...")
    build_slide_3(prs)
    print("Building Slide 4: Problem Statement & Research Gap...")
    build_slide_4(prs)
    print("Building Slide 5: Proposed Methodology / MSBA-Net (Only Diagram Slide)...")
    build_slide_5(prs)
    print("Building Slide 6: Experimental Work: Dataset & Setup...")
    build_slide_6(prs)
    print("Building Slide 7: Experimental Work: Training & Evaluation...")
    build_slide_7(prs)
    print("Building Slide 8: Results: Comparison with Existing Methods...")
    build_slide_8(prs)
    print("Building Slide 9: Results: Detailed Analysis...")
    build_slide_9(prs)
    print("Building Slide 10: Ablation Study & Analysis...")
    build_slide_10(prs)
    print("Building Slide 11: Limitations & Robustness Analysis...")
    build_slide_11(prs)
    print("Building Slide 12: Conclusions & Future Work...")
    build_slide_12(prs)
    print("Building Slide 13: Thank You...")
    build_slide_13(prs)

    output_path = os.path.join(os.path.dirname(__file__), "MSBA-Net_Amen_Parmar_Presentation.pptx")
    prs.save(output_path)
    print(f"SUCCESS: Generated {output_path} with EXACTLY {len(prs.slides)} slides!")

if __name__ == "__main__":
    main()
